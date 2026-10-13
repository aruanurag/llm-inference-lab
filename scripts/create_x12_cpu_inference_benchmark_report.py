from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/x12-long-prefill-long-decode-benchmark-report.docx")

NAVY = "17365D"
BLUE = "DCEAF7"
PALE_BLUE = "F3F7FB"
PALE_GREEN = "EAF5EF"
PALE_GOLD = "FFF6DD"
PALE_GRAY = "F7F8FA"
GRID = "D7DEE7"
BLACK = RGBColor(0, 0, 0)
MUTED = RGBColor(78, 90, 104)
WHITE = RGBColor(255, 255, 255)
BLUE_TEXT = RGBColor(23, 54, 93)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = GRID) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_column_width(cell, width_inches: float) -> None:
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    color: RGBColor = BLACK,
    size: float = 8.7,
    align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    set_cell_borders(cell)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    *,
    font_size: float = 8.5,
    header_fill: str = NAVY,
    alternate_fill: str = PALE_BLUE,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_column_width(cell, widths[i])
        set_cell_shading(cell, header_fill)
        set_cell_text(
            cell,
            header,
            bold=True,
            color=WHITE,
            size=max(font_size - 0.1, 7.8),
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        fill = "FFFFFF" if row_index % 2 == 0 else alternate_fill
        for i, value in enumerate(row):
            set_column_width(cells[i], widths[i])
            set_cell_shading(cells[i], fill)
            align = WD_ALIGN_PARAGRAPH.CENTER if i > 0 and len(headers) >= 5 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[i], value, size=font_size, align=align)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_callout(doc: Document, title: str, text: str, fill: str = PALE_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, start=180, bottom=130, end=180)
    set_cell_borders(cell, "B7C9DD")
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Aptos"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    r.font.size = Pt(10)
    r.font.color.rgb = BLUE_TEXT
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.08
    r2 = p2.add_run(text)
    r2.font.name = "Aptos"
    r2._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    r2._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    r2.font.size = Pt(9)
    r2.font.color.rgb = BLACK
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    p.add_run(text)


def add_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(8.6)
    run.font.color.rgb = MUTED


def remove_paragraph_borders(paragraph) -> None:
    """Remove inherited Office title/subtitle rules for a clean report cover."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, face, size, before, after in (
        ("Title", "Aptos Display", 25, 0, 8),
        ("Subtitle", "Aptos", 12, 0, 12),
        ("Heading 1", "Aptos", 15, 15, 6),
        ("Heading 2", "Aptos", 11.5, 10, 4),
    ):
        style = styles[style_name]
        style.font.name = face
        style._element.rPr.rFonts.set(qn("w:ascii"), face)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), face)
        style.font.size = Pt(size)
        style.font.bold = style_name != "Subtitle"
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run("OCI CPU Inference Benchmark Report  |  X12 long prefill and long decode")
    r.font.name = "Aptos"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED


def main() -> None:
    doc = Document()
    style_document(doc)
    doc.core_properties.title = "X12 CPU Inference Benchmark Report"
    doc.core_properties.subject = "Long prefill and long decode measurements on OCI CPU"
    doc.core_properties.author = "OCI CPU Inference Lab"

    # Cover / executive summary.
    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title_run = title.add_run("X12 CPU Inference Benchmark Report")
    title_run.font.name = "Aptos Display"
    title_run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title_run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title_run.font.size = Pt(25)
    title_run.font.bold = True
    title_run.font.color.rgb = BLACK
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Long-prompt prefill and long-output decode on OCI CPU with llama.cpp")
    doc.add_paragraph(
        "This report documents a single-shape CPU inference experiment. It measures end-to-end streaming behavior under increasing client concurrency, and records native llama-bench reference measurements separately. It is designed to be shareable without infrastructure identifiers, endpoints, credentials, or raw local paths."
    )
    add_callout(
        doc,
        "Executive readout",
        "With four active server slots, client concurrency of 4 was the practical balance point in these runs. Above 4 concurrent clients, aggregate output rate remained broadly flat while p95 time to first token increased materially because requests waited for an available server slot. This is a capacity observation for this exact configuration—not a general claim about all OCI CPU shapes or models.",
        PALE_GREEN,
    )
    doc.add_heading("Test at a glance", level=1)
    add_table(
        doc,
        ["Area", "Measured configuration"],
        [
            ["Platform", "OCI VM.Standard4.Ax.Flex; 8 OCPUs; 64 GiB memory"],
            ["Model", "Qwen2.5 1.5B Instruct, Q4_K_M GGUF quantization"],
            ["Serving engine", "llama.cpp server, native build enabled"],
            ["Workloads", "Long prompt prefill; short prompt / long output decode"],
            ["Serving tests", "HTTP streaming; 50 requests per run; client concurrency 1, 2, 4, 8, and 16"],
            ["Native reference", "llama-bench; 50 repetitions; isolated prompt and decode paths"],
        ],
        [1.45, 5.65],
        font_size=9,
    )
    add_note(
        doc,
        "Scope: one model, one engine build, one server configuration, and one OCI VM shape. Results should be repeated with the same configuration before making comparative hardware conclusions."
    )

    doc.add_page_break()
    doc.add_heading("1. Hardware, model, and serving configuration", level=1)
    doc.add_paragraph(
        "The same platform, model, and llama.cpp configuration were used for both workload families. This makes the tables comparable within this report; it does not make them comparable to a different shape, model quantization, engine version, or server tuning profile."
    )
    add_table(
        doc,
        ["Layer", "Configuration"],
        [
            ["Compute", "OCI VM.Standard4.Ax.Flex; 8 OCPUs; 64 GiB memory"],
            ["Operating mode", "One llama.cpp server process on the VM"],
            ["Inference engine", "llama.cpp / GGUF serving"],
            ["Model", "Qwen2.5 1.5B Instruct Q4_K_M"],
            ["Context window", "4,096 tokens"],
            ["Generation CPU threads", "8"],
            ["Prompt / prefill CPU threads", "8"],
            ["Parallel slots", "4 active server-side request slots"],
            ["Batch / micro-batch", "512 / 128"],
            ["Build profile", "Native build enabled; VNNI option disabled in this deployment profile"],
        ],
        [2.05, 5.05],
        font_size=9,
    )
    add_callout(
        doc,
        "Why the four parallel slots matter",
        "Client concurrency is how many requests the benchmark sends at once. Parallel slots are how many requests the server can actively hold. At concurrency 8 or 16, this server configuration can queue work behind four active slots. Higher p95 TTFT therefore includes queueing time as well as prompt processing time.",
        PALE_GOLD,
    )

    doc.add_heading("2. Workloads and prompts", level=1)
    doc.add_paragraph(
        "Both test families rotate through three prompts. The input-token figure below is a planning estimate used by the lab's native benchmark path (prompt characters divided by four and rounded); it is not an exact tokenizer count. Output is controlled with an upper generation limit, not a guarantee that every response contains exactly that many tokens."
    )
    add_table(
        doc,
        ["Workload", "Prompt set and real-world analogue", "Inputs", "Output limit", "Primary question"],
        [
            [
                "Long prompt prefill",
                "Three technical prompts about local LLM serving, CPU tuning, batching, queue delay, and prefill/decode behavior. Similar to asking an assistant to analyze a detailed technical brief.",
                "3 prompts; ~70 planning tokens per prompt",
                "128 tokens",
                "How quickly does the server begin streaming after processing a longer input?",
            ],
            [
                "Long output decode",
                "Three short prompts requesting detailed checklists, a llama.cpp tuning guide, and a CPU-vs-GPU comparison. Similar to creative or explanatory requests with a short instruction and long response.",
                "3 prompts; ~20 planning tokens per prompt",
                "512 tokens",
                "How does sustained generation behave as concurrent clients increase?",
            ],
        ],
        [1.15, 2.55, 1.05, 0.8, 1.55],
        font_size=8.1,
    )
    doc.add_heading("Prompt inventory", level=2)
    add_table(
        doc,
        ["Set", "Prompt themes used"],
        [
            ["Long prompt prefill", "(1) llama.cpp, GGUF, quantization, batching, and CPU/GPU serving; (2) a mixed short/long CPU-serving scenario with rising TTFT; (3) a detailed explanation of prefill, decode, queue delay, concurrency, and p95 behavior."],
            ["Long output decode", "(1) detailed checklist for benchmarking a local LLM server; (2) step-by-step guide for tuning llama.cpp on a CPU instance; (3) detailed comparison of CPU and GPU inference for small language models."],
        ],
        [1.45, 5.65],
        font_size=8.5,
    )

    doc.add_heading("3. HTTP streaming results — long prompt prefill", level=1)
    doc.add_paragraph(
        "Each row is a completed HTTP streaming run with 50 requests, all successful, prompt-cache reuse disabled, and the 128-token output cap. For repeated concurrency levels, the latest completed run is shown. TTFT, inter-token latency (ITL), and end-to-end latency are p95 values."
    )
    add_table(
        doc,
        ["Client\nconcurrency", "Success", "TTFT\np95", "Approx. ITL\np95", "E2E latency\np95", "Request\nrate", "Approx. output\nrate"],
        [
            ["1", "50 / 50", "333 ms", "23 ms", "2.68 s", "0.380 req/s", "60.8 tok/s"],
            ["2", "50 / 50", "584 ms", "45 ms", "5.31 s", "0.475 req/s", "77.4 tok/s"],
            ["4", "50 / 50", "998 ms", "46 ms", "4.67 s", "0.847 req/s", "136.5 tok/s"],
            ["8", "50 / 50", "5.53 s", "46 ms", "9.19 s", "0.835 req/s", "135.1 tok/s"],
            ["16", "50 / 50", "14.57 s", "45 ms", "18.25 s", "0.857 req/s", "138.7 tok/s"],
        ],
        [0.78, 0.77, 0.85, 0.95, 0.95, 0.88, 1.02],
        font_size=8.2,
    )
    add_callout(
        doc,
        "Interpretation for this workload",
        "Throughput rose from 60.8 to 136.5 approximate output tokens/s when client concurrency increased from 1 to 4. The 4-client run was the highest-throughput point before queueing became visible. At 8 and 16 clients, throughput stayed near 135–139 approximate output tokens/s, but p95 TTFT rose to 5.53 s and 14.57 s respectively. The flat ITL near 45–46 ms suggests that the later slowdown is dominated by waiting for service capacity, rather than a collapse in observed streaming cadence once a response starts.",
    )
    add_note(
        doc,
        "The concurrency-2 repeat was stable: 76.9 and 77.4 approximate output tokens/s, with p95 TTFT of 575 ms and 584 ms. The concurrency-16 repeat was also stable: both runs produced about 138.7 approximate output tokens/s, with p95 TTFT of 14.62 s and 14.57 s."
    )

    doc.add_heading("4. HTTP streaming results — long output decode", level=1)
    doc.add_paragraph(
        "Each row is a completed HTTP streaming run with 50 requests, all successful, prompt-cache reuse disabled, and the 512-token output cap. The later of two completed concurrency-8 runs is shown."
    )
    add_table(
        doc,
        ["Client\nconcurrency", "Success", "TTFT\np95", "Approx. ITL\np95", "E2E latency\np95", "Request\nrate", "Approx. output\nrate"],
        [
            ["1", "50 / 50", "200 ms", "23 ms", "9.85 s", "0.108 req/s", "62.9 tok/s"],
            ["2", "50 / 50", "346 ms", "45 ms", "19.42 s", "0.122 req/s", "73.1 tok/s"],
            ["4", "50 / 50", "574 ms", "48 ms", "16.38 s", "0.252 req/s", "149.9 tok/s"],
            ["8", "50 / 50", "16.37 s", "47 ms", "32.21 s", "0.251 req/s", "152.2 tok/s"],
            ["16", "50 / 50", "48.16 s", "47 ms", "63.83 s", "0.249 req/s", "151.9 tok/s"],
        ],
        [0.78, 0.77, 0.85, 0.95, 0.95, 0.88, 1.02],
        font_size=8.2,
    )
    add_callout(
        doc,
        "Interpretation for this workload",
        "Aggregate output rate reached roughly 150 approximate tokens/s at concurrency 4 and remained around that level at 8 and 16. However, the p95 time to first token changed from 574 ms at concurrency 4 to 16.37 s at 8 and 48.16 s at 16. For an interactive service using this four-slot configuration, increasing client concurrency past 4 trades responsiveness for queued throughput rather than delivering meaningful extra aggregate generation capacity.",
        PALE_GOLD,
    )

    doc.add_heading("5. Native llama-bench reference measurements", level=1)
    doc.add_paragraph(
        "The same experiment also ran llama-bench. This is an engine-native microbenchmark: it repeats isolated prompt and generation work locally and does not create network requests. It therefore reports prompt and decode token rates, but not HTTP TTFT, request rate, or streamed ITL. These values are a complementary engine reference and should not be directly ranked against the end-to-end HTTP tables."
    )
    add_table(
        doc,
        ["Workload", "Repetitions", "Planning input", "Output limit", "Prompt rate", "Decode rate", "Interpretation"],
        [
            ["Long prompt prefill (C1)", "50", "70 tokens", "128", "359.2 tok/s", "55.6 tok/s", "Isolated local prompt/decode reference"],
            ["Long prompt prefill (C2)", "50", "70 tokens", "128", "358.8 tok/s", "55.6 tok/s", "Repeat at metadata concurrency 2"],
            ["Long prompt prefill (C4)", "50", "70 tokens", "128", "350.1 tok/s", "55.5 tok/s", "Repeat at metadata concurrency 4"],
            ["Long output decode (C4)", "50", "20 tokens", "512", "299.9 tok/s", "56.2 tok/s", "Long-generation reference"],
            ["Long output decode (C8)", "50", "20 tokens", "512", "307.2 tok/s", "56.2 tok/s", "Repeat at metadata concurrency 8"],
        ],
        [1.55, 0.68, 0.8, 0.65, 0.8, 0.8, 1.02],
        font_size=7.9,
    )
    add_note(
        doc,
        "For llama-bench, the C label is retained from experiment metadata. It should not be read as multi-client HTTP concurrency; the tool is a local microbenchmark and its principal repetition control was 50."
    )

    doc.add_heading("6. Measurement definitions and caveats", level=1)
    add_table(
        doc,
        ["Metric", "Meaning in this report"],
        [
            ["TTFT p95", "95th-percentile time from HTTP request start until the first non-empty streamed text chunk. It contains queue time when more client requests arrive than the four configured server slots can serve."],
            ["Approx. ITL p95", "95th-percentile gap between non-empty streamed chunks. It is an SSE chunk-gap estimate, not a tokenizer-level per-token measurement."],
            ["E2E latency p95", "95th-percentile time from HTTP request start to stream completion."],
            ["Request rate", "Completed HTTP requests divided by wall-clock time for the full run."],
            ["Approx. output rate", "Generated response characters divided by four and by wall-clock time. This is a directional approximation, not an exact tokenized output rate."],
            ["Native prompt/decode rate", "llama-bench local microbenchmark rate; excludes network, queuing, request parsing, and streaming behavior."],
        ],
        [1.55, 5.55],
        font_size=8.6,
    )
    doc.add_heading("Conclusions and next measurements", level=1)
    add_bullet(doc, "For the tested Qwen2.5 1.5B Q4_K_M / llama.cpp configuration, four concurrent HTTP clients were the practical interactive capacity point because it matched the server's four active slots.")
    add_bullet(doc, "Long-prefill and long-decode throughput did not increase materially beyond four clients, while p95 TTFT increased sharply as queued work accumulated.")
    add_bullet(doc, "The native llama-bench references were stable across their recorded metadata settings, at roughly 350–359 prompt tokens/s for the prefill set and about 55–56 decode tokens/s for this model and engine profile.")
    add_bullet(doc, "Next: repeat on a second shape with the same model, engine version, thread counts, slot count, prompt sets, cache state, and request schedule; then separately increase parallel slots and CPU threads to map the latency-throughput tradeoff.")
    doc.add_heading("What this report does—and does not—establish", level=2)
    add_table(
        doc,
        ["Supported by these runs", "Not established by these runs"],
        [
            ["The latency-throughput behavior of this single VM, model, engine, and four-slot server configuration.", "A general performance conclusion about other OCI CPU shapes, model sizes, or llama.cpp versions."],
            ["That work began queueing when HTTP client concurrency exceeded the configured four active slots.", "A tokenizer-exact output-token rate or true per-token latency; the HTTP values are explicitly approximated."],
            ["That native prompt/decode measurements and end-to-end serving metrics answer different questions.", "The benefit of prompt caching, because cache reuse was disabled for the HTTP test runs."],
        ],
        [3.55, 3.55],
        font_size=8.4,
    )
    doc.add_heading("Next comparison protocol", level=2)
    add_table(
        doc,
        ["Keep fixed", "Change deliberately", "Report"],
        [
            ["Model, GGUF quantization, llama.cpp build/version, 8 OCPUs, memory per OCPU, context, batch, micro-batch, threads, slots, prompt cache state, prompt sets, and requests per run.", "Only the OCI shape for a hardware comparison; then, in separate tests, parallel slots and thread counts for a server-tuning comparison.", "Median and p95 TTFT, p95 approximate ITL, p95 end-to-end latency, request rate, approximate output rate, success count, and the native llama-bench prompt/decode reference."],
        ],
        [2.45, 2.15, 2.5],
        font_size=8.1,
    )
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
